"""Document parser supporting PDF, DOCX, TXT, Markdown, and CSV.

Validates file signatures, extensions, sizes, and non-empty extracted text.
Uses pypdf / python-docx when installed and includes standard-library PDF stream
and DOCX XML (`word/document.xml`) parsers so ingestion works in all environments.
"""

import io
import re
import xml.etree.ElementTree as ET
import zipfile
import zlib
from typing import Optional
from app.core.exceptions import ValidationError
from app.utils.csv_parser import parse_csv_to_structured_text
from app.utils.validators import validate_upload_metadata


class DocumentParserService:
    """Extracts clean plain text from PDF, DOCX, TXT, Markdown, and CSV documents."""

    def __init__(self, max_upload_size_mb: int = 10) -> None:
        self.max_upload_size_mb = max_upload_size_mb

    def parse_document(
        self,
        filename: str,
        content: bytes,
        content_type: str = "application/octet-stream",
    ) -> str:
        """Validate and extract text from an uploaded document."""
        ext = validate_upload_metadata(
            filename=filename,
            content_type=content_type,
            size_bytes=len(content),
            max_mb=self.max_upload_size_mb,
        )

        if ext in (".txt", ".md", ".markdown"):
            extracted = self._parse_text(content)
        elif ext == ".csv":
            extracted = parse_csv_to_structured_text(content)
        elif ext == ".docx":
            extracted = self._parse_docx(content)
        elif ext == ".pdf":
            extracted = self._parse_pdf(content)
        else:
            raise ValidationError(f"Unsupported file extension '{ext}'.")

        cleaned = re.sub(r"\n{3,}", "\n\n", (extracted or "").strip())
        if not cleaned:
            raise ValidationError(f"Document '{filename}' contains no extractable text content.")
        return cleaned

    def _parse_text(self, content: bytes) -> str:
        if content.startswith(b"\x00\x00") or b"\x00" in content[:256]:
            raise ValidationError("Binary file content detected in text document.")
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="replace")

    def _parse_docx(self, content: bytes) -> str:
        if not content.startswith(b"PK\x03\x04"):
            raise ValidationError("Invalid DOCX file signature (expected ZIP container).")

        # Try python-docx if installed
        try:
            import docx  # type: ignore

            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            if paragraphs:
                return "\n\n".join(paragraphs)
        except ImportError:
            pass
        except Exception as exc:
            raise ValidationError(f"Failed to parse DOCX document: {exc}") from exc

        # Standard library OOXML parser (`word/document.xml`)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                if "word/document.xml" not in zf.namelist():
                    raise ValidationError("Invalid DOCX archive: missing word/document.xml.")
                xml_bytes = zf.read("word/document.xml")
            root = ET.fromstring(xml_bytes)
            paragraphs = []
            for p_elem in root.iter():
                if p_elem.tag.endswith("}p") or p_elem.tag == "p":
                    texts = [
                        t_elem.text
                        for t_elem in p_elem.iter()
                        if (t_elem.tag.endswith("}t") or t_elem.tag == "t") and t_elem.text
                    ]
                    line = "".join(texts).strip()
                    if line:
                        paragraphs.append(line)
            return "\n\n".join(paragraphs)
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError(f"Corrupted DOCX document: {exc}") from exc

    def _parse_pdf(self, content: bytes) -> str:
        if not content.lstrip().startswith(b"%PDF-"):
            raise ValidationError("Invalid PDF file signature (missing %PDF- header).")

        # Try pypdf if installed
        try:
            import pypdf  # type: ignore

            reader = pypdf.PdfReader(io.BytesIO(content))
            pages = []
            for page in reader.pages:
                t = page.extract_text() or ""
                if t.strip():
                    pages.append(t.strip())
            if pages:
                return "\n\n".join(pages)
        except ImportError:
            pass
        except Exception as exc:
            raise ValidationError(f"Failed to parse PDF document: {exc}") from exc

        # Standard-library PDF content stream & literal string extractor
        extracted_fragments = []
        streams = re.findall(rb"stream\r?\n(.*?)\r?\nendstream", content, re.DOTALL)
        buffers = [content]
        for s in streams:
            buffers.append(s)
            try:
                buffers.append(zlib.decompress(s))
            except Exception:
                pass

        for buf in buffers:
            # Extract PDF text literals inside parentheses: (text) Tj / TJ
            for match in re.finditer(rb"\(([^()\\]*(?:\\.[^()\\]*)*)\)", buf):
                raw_bytes = match.group(1)
                raw_bytes = (
                    raw_bytes.replace(rb"\(", b"(")
                    .replace(rb"\)", b")")
                    .replace(rb"\n", b"\n")
                    .replace(rb"\r", b" ")
                    .replace(rb"\\", b"\\")
                )
                try:
                    decoded = raw_bytes.decode("utf-8").strip()
                except UnicodeDecodeError:
                    decoded = raw_bytes.decode("latin-1", errors="ignore").strip()
                if len(decoded) >= 2 and re.search(r"[A-Za-z0-9]", decoded):
                    extracted_fragments.append(decoded)

        deduped = []
        for frag in extracted_fragments:
            if not deduped or deduped[-1] != frag:
                deduped.append(frag)
        return "\n".join(deduped)
